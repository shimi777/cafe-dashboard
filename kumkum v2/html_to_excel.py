import re
from datetime import datetime
from bs4 import BeautifulSoup
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import io

def parse_html_transactions(html_content):
    """
    Parse transaction details from HTML
    Returns list of transactions with their details
    """
    soup = BeautifulSoup(html_content, 'html.parser')
    
    transactions = []
    
    # Find all data blocks (each block = one transaction)
    data_blocks = soup.find_all('div', class_='data-block')
    
    for block in data_blocks:
        # Extract transaction header
        trans_header = block.find('div', class_='trans-header')
        if not trans_header:
            continue
            
        # Get transaction details from header
        header_items = trans_header.find_all('span', class_='header-num')
        
        if len(header_items) < 4:
            continue
        
        order_id = header_items[0].get_text(strip=True)
        invoice_num = header_items[1].get_text(strip=True)
        transaction_type = header_items[2].get_text(strip=True)
        register_num = header_items[3].get_text(strip=True)
        datetime_str = header_items[4].get_text(strip=True)
        
        # Parse date and time
        try:
            trans_datetime = datetime.strptime(datetime_str, '%d/%m/%Y %H:%M')
            trans_date = trans_datetime.date()
            trans_time = trans_datetime.time()
        except:
            continue
        
        # Get items
        items = []
        item_rows = block.find_all('div', class_='item-row')
        
        total_amount = 0
        
        for item_row in item_rows:
            # Skip tender and totals rows
            if 'tender-row' in item_row.get('class', []) or 'totals-row' in item_row.get('class', []):
                continue
            
            # Get item description
            item_name = item_row.find('span')
            if not item_name:
                continue
            item_name = item_name.get_text(strip=True)
            
            # Get prices from num elements
            num_elements = item_row.find_all('div', class_='num')
            
            if len(num_elements) >= 4:
                try:
                    quantity = float(num_elements[0].get_text(strip=True))
                    unit_price = float(num_elements[1].get_text(strip=True))
                    taxable_price = float(num_elements[2].get_text(strip=True))
                    total_price = float(num_elements[3].get_text(strip=True))
                    
                    items.append({
                        'name': item_name,
                        'quantity': quantity,
                        'unit_price': unit_price,
                        'taxable_price': taxable_price,
                        'total_price': total_price
                    })
                    
                    total_amount += total_price
                except (ValueError, IndexError):
                    continue
        
        # Get total from totals row
        totals_row = block.find('div', class_='totals-row')
        if totals_row:
            total_nums = totals_row.find_all('span', class_='tender-num')
            if len(total_nums) >= 3:
                try:
                    total_amount = float(total_nums[2].get_text(strip=True))
                except:
                    pass
        
        # Skip transactions with total = 0 (cancelled transactions)
        if total_amount == 0:
            continue
        
        transactions.append({
            'order_id': order_id,
            'invoice_num': invoice_num,
            'date': trans_date,
            'time': trans_time,
            'items': items,
            'total': total_amount
        })
    
    return transactions

def create_daily_summary(transactions):
    """
    Summarize transactions by date
    Returns dataframe with daily sales summary
    """
    daily_data = {}
    
    for trans in transactions:
        trans_date = trans['date']
        total = trans['total']
        
        if trans_date not in daily_data:
            daily_data[trans_date] = {
                'date': trans_date,
                'total_sales': 0,
                'transaction_count': 0,
                'items_count': 0
            }
        
        daily_data[trans_date]['total_sales'] += total
        daily_data[trans_date]['transaction_count'] += 1
        daily_data[trans_date]['items_count'] += len(trans['items'])
    
    # Convert to dataframe
    daily_df = pd.DataFrame(list(daily_data.values()))
    daily_df = daily_df.sort_values('date')
    
    return daily_df

def create_detailed_transactions_df(transactions):
    """
    Create detailed transactions dataframe (one row per transaction)
    """
    records = []
    
    for trans in transactions:
        records.append({
            'Order ID': trans['order_id'],
            'Invoice': trans['invoice_num'],
            'Date': trans['date'],
            'Time': trans['time'],
            'Item Count': len(trans['items']),
            'Total Amount': trans['total']
        })
    
    return pd.DataFrame(records)

def create_items_summary_df(transactions):
    """
    Create item-level summary (aggregated by item name)
    """
    items_data = {}
    
    for trans in transactions:
        for item in trans['items']:
            item_name = item['name']
            
            if item_name not in items_data:
                items_data[item_name] = {
                    'item_name': item_name,
                    'quantity': 0,
                    'total_amount': 0,
                    'transaction_count': 0
                }
            
            items_data[item_name]['quantity'] += item['quantity']
            items_data[item_name]['total_amount'] += item['total_price']
            items_data[item_name]['transaction_count'] += 1
    
    items_df = pd.DataFrame(list(items_data.values()))
    items_df = items_df.sort_values('total_amount', ascending=False)
    
    return items_df

def export_to_excel(transactions, filepath):
    """
    Export all data to Excel file with multiple sheets
    """
    wb = Workbook()
    
    # Daily Summary Sheet
    daily_df = create_daily_summary(transactions)
    ws_daily = wb.active
    ws_daily.title = "Daily Summary"
    
    # Add headers
    headers = ['Date', 'Total Sales', 'Transaction Count', 'Items Count']
    for col_idx, header in enumerate(headers, 1):
        ws_daily.cell(row=1, column=col_idx, value=header)
        ws_daily.cell(row=1, column=col_idx).font = Font(bold=True)
        ws_daily.cell(row=1, column=col_idx).fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        ws_daily.cell(row=1, column=col_idx).font = Font(bold=True, color="FFFFFF")
    
    # Add data
    for row_idx, row in enumerate(daily_df.values, 2):
        ws_daily.cell(row=row_idx, column=1, value=row[0])  # date
        ws_daily.cell(row=row_idx, column=2, value=row[1])  # total_sales
        ws_daily.cell(row=row_idx, column=3, value=row[2])  # transaction_count
        ws_daily.cell(row=row_idx, column=4, value=row[3])  # items_count
    
    ws_daily.column_dimensions['A'].width = 15
    ws_daily.column_dimensions['B'].width = 15
    ws_daily.column_dimensions['C'].width = 18
    ws_daily.column_dimensions['D'].width = 15
    
    # Detailed Transactions Sheet
    trans_df = create_detailed_transactions_df(transactions)
    ws_trans = wb.create_sheet("Transactions")
    
    headers = ['Order ID', 'Invoice', 'Date', 'Time', 'Item Count', 'Total Amount']
    for col_idx, header in enumerate(headers, 1):
        ws_trans.cell(row=1, column=col_idx, value=header)
        ws_trans.cell(row=1, column=col_idx).font = Font(bold=True, color="FFFFFF")
        ws_trans.cell(row=1, column=col_idx).fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    
    for row_idx, row in enumerate(trans_df.values, 2):
        ws_trans.cell(row=row_idx, column=1, value=row[0])  # order_id
        ws_trans.cell(row=row_idx, column=2, value=row[1])  # invoice
        ws_trans.cell(row=row_idx, column=3, value=row[2])  # date
        ws_trans.cell(row=row_idx, column=4, value=row[3])  # time
        ws_trans.cell(row=row_idx, column=5, value=row[4])  # items
        ws_trans.cell(row=row_idx, column=6, value=row[5])  # total
    
    ws_trans.column_dimensions['A'].width = 12
    ws_trans.column_dimensions['B'].width = 12
    ws_trans.column_dimensions['C'].width = 12
    ws_trans.column_dimensions['D'].width = 12
    ws_trans.column_dimensions['E'].width = 12
    ws_trans.column_dimensions['F'].width = 15
    
    # Items Summary Sheet
    items_df = create_items_summary_df(transactions)
    ws_items = wb.create_sheet("Items Summary")
    
    headers = ['Item Name', 'Quantity', 'Total Amount', 'Transaction Count']
    for col_idx, header in enumerate(headers, 1):
        ws_items.cell(row=1, column=col_idx, value=header)
        ws_items.cell(row=1, column=col_idx).font = Font(bold=True, color="FFFFFF")
        ws_items.cell(row=1, column=col_idx).fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    
    for row_idx, row in enumerate(items_df.values, 2):
        ws_items.cell(row=row_idx, column=1, value=row[0])  # name
        ws_items.cell(row=row_idx, column=2, value=row[1])  # quantity
        ws_items.cell(row=row_idx, column=3, value=row[2])  # total
        ws_items.cell(row=row_idx, column=4, value=row[3])  # count
    
    ws_items.column_dimensions['A'].width = 25
    ws_items.column_dimensions['B'].width = 12
    ws_items.column_dimensions['C'].width = 15
    ws_items.column_dimensions['D'].width = 18
    
    # Save
    wb.save(filepath)
    
    return daily_df, trans_df, items_df
